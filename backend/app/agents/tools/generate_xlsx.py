"""generate_plan_xlsx 工具：根据营销方案结构化数据生成 XLSX 预算流程回报分析表格。

LangChain @tool 格式，与 web_search / web_fetch 同一模式。
供 plan_generator_agent handler 内部直接调用，也可被 ReAct agent 绑定使用。

参考格式：娃哈哈营销方案_预算流程回报分析.xlsx
"""

import json
import logging
from datetime import date
from pathlib import Path

import openpyxl
from langchain_core.tools import tool
from openpyxl.chart import BarChart, PieChart, LineChart, RadarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field

from app.schemas.xlsx_generation import XlsxData

logger = logging.getLogger(__name__)

# ── 输出目录 ──
OUTPUT_DIR = Path("generated_xlsx")

# ── 样式常量 ──
TITLE_FONT = Font(name="微软雅黑", size=16, bold=True, color="FFFFFF")
TITLE_FILL = PatternFill("solid", fgColor="1F4E79")
HEAD_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="2D3748")
SUBHEAD_FONT = Font(name="微软雅黑", size=11, bold=True, color="1F4E79")
SUBHEAD_FILL = PatternFill("solid", fgColor="D6E4F0")
CELL_FONT = Font(name="微软雅黑", size=10, color="2D3748")
TOTAL_FONT = Font(name="微软雅黑", size=11, bold=True, color="1F4E79")
TOTAL_FILL = PatternFill("solid", fgColor="FFF2CC")
GREEN_FILL = PatternFill("solid", fgColor="E2EFDA")
ORANGE_FILL = PatternFill("solid", fgColor="FCE4D6")
BLUE_FILL = PatternFill("solid", fgColor="DDEBF7")
GREY_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN_BORDER = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

# 阶段颜色映射
PHASE_FILLS = {
    "筹备期": PatternFill("solid", fgColor="D6E4F0"),
    "预热期": PatternFill("solid", fgColor="FFF2CC"),
    "爆发期": PatternFill("solid", fgColor="FCE4D6"),
    "收割期": PatternFill("solid", fgColor="E2EFDA"),
}

# 维度颜色映射
DIM_FILLS = {
    "品牌传播": PatternFill("solid", fgColor="D6E4F0"),
    "用户增长": PatternFill("solid", fgColor="E2EFDA"),
    "销售转化": PatternFill("solid", fgColor="FCE4D6"),
    "渠道建设": PatternFill("solid", fgColor="E4DFEC"),
}

# 预算颜色映射
CAT_FILLS: dict[str, PatternFill] = {}


def _get_cat_fill(cat: str) -> PatternFill:
    if cat not in CAT_FILLS:
        # 循环使用一组颜色
        colors = [
            "E2EFDA", "DDEBF7", "FCE4D6", "E4DFEC",
            "FFF2CC", "FCE4D6", "D9E1F2", "F2F2F2",
        ]
        idx = len(CAT_FILLS) % len(colors)
        CAT_FILLS[cat] = PatternFill("solid", fgColor=colors[idx])
    return CAT_FILLS[cat]


class GeneratePlanXlsxInput(BaseModel):
    """generate_plan_xlsx 工具输入。"""

    data: str = Field(
        ..., description="XlsxData 的完整 JSON 字符串（由 LLM 结构化输出序列化而来）"
    )
    brand_name: str = Field(..., description="品牌名称，用于文件名")


def _set_widths(ws, widths: list[float]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _apply_cell(
    cell,
    value,
    font=CELL_FONT,
    fill=None,
    align=CENTER,
    border=True,
    num_fmt=None,
) -> None:
    cell.value = value
    cell.font = font
    if fill:
        cell.fill = fill
    if border:
        cell.border = THIN_BORDER
    cell.alignment = align
    if num_fmt:
        cell.number_format = num_fmt


# ── Sheet 1: 预算总览对比 ──────────────────────────────────


def _build_sheet1(wb, data: XlsxData) -> None:
    bi = data.brand_info
    bo = data.budget_overview
    ws = wb.active
    has_vb = bi.version_b_name is not None
    ws.title = "预算总览"

    # 标题行
    title_text = f"{bi.brand_name}营销方案 — 预算总览"
    col_count = 7 if has_vb else 4
    last_col_letter = get_column_letter(col_count)
    ws.merge_cells(f"A1:{last_col_letter}1")
    _apply_cell(
        ws["A1"], title_text,
        TITLE_FONT, TITLE_FILL, CENTER,
    )
    ws.row_dimensions[1].height = 36

    ws.merge_cells(f"A2:{last_col_letter}2")
    _apply_cell(
        ws["A2"], bi.subtitle,
        Font(name="微软雅黑", size=10, color="718096"), align=CENTER,
    )
    ws.row_dimensions[2].height = 22

    # 表头
    va, vb = bi.version_a_name, bi.version_b_name
    if has_vb:
        headers = [
            "序号", "费用项目",
            f"{va}金额(万元)", f"{va}占比",
            f"{vb}金额(万元)", f"{vb}占比",
            "差异说明",
        ]
    else:
        headers = ["序号", "费用项目", "金额(万元)", "占比"]

    for i, h in enumerate(headers, 1):
        _apply_cell(ws.cell(row=3, column=i), h, HEAD_FONT, HEAD_FILL)
    ws.row_dimensions[3].height = 28

    # 数据行
    for idx, row in enumerate(bo.rows, 4):
        if has_vb:
            row_vals = [
                idx - 3,
                row.category,
                row.amount_version_a,
                row.pct_version_a,
                row.amount_version_b,
                row.pct_version_b,
                row.diff,
            ]
        else:
            row_vals = [
                idx - 3,
                row.category,
                row.amount_version_a,
                row.pct_version_a,
            ]

        row_fill = GREEN_FILL if (idx - 3) % 2 == 1 else None

        for c, val in enumerate(row_vals, 1):
            align = LEFT if c in (2, col_count) else CENTER
            num_fmt = "0.0" if c in (3, 5) else None
            _apply_cell(
                ws.cell(row=idx, column=c), val,
                CELL_FONT, row_fill, align,
                num_fmt=num_fmt,
            )
        ws.row_dimensions[idx].height = 26

    # 合计行
    total_row = 4 + len(bo.rows)
    if has_vb:
        total_vals = ["", "合计", bo.total_a, "100%", bo.total_b, "100%"]
        diff_total = (bo.total_a - bo.total_b) if bo.total_b is not None else 0
        if diff_total != 0:
            total_vals.append(f"{va}比{vb}多{diff_total}万")
        else:
            total_vals.append(f"{va}与{vb}预算一致")
    else:
        total_vals = ["", "合计", bo.total_a, "100%"]

    for c, val in enumerate(total_vals, 1):
        align = LEFT if c in (2, col_count) else CENTER
        num_fmt = "0.0" if c in (3, 5) else None
        _apply_cell(
            ws.cell(row=total_row, column=c), val,
            TOTAL_FONT, TOTAL_FILL, align, num_fmt=num_fmt,
        )
    ws.row_dimensions[total_row].height = 30

    # 备注行
    note_row = total_row + 2
    ws.merge_cells(f"A{note_row}:{last_col_letter}{note_row}")
    _apply_cell(
        ws.cell(row=note_row, column=1),
        f"备注：{bo.note}",
        Font(name="微软雅黑", size=9, color="718096", italic=True),
        align=LEFT,
    )

    col_widths = [6, 28, 16, 12]
    if has_vb:
        col_widths += [16, 12, 34]
    _set_widths(ws, col_widths)

    # 图表（单版本显示饼图，双版本显示柱状图+饼图）
    if len(bo.rows) > 1:
        if has_vb:
            _add_charts_sheet1(ws, bo, va, vb, total_row)
        else:
            _add_single_pie_chart(ws, bo, va, total_row)


def _add_single_pie_chart(ws, bo, va, total_row) -> None:
    """单版本预算总览表：只加一个饼图。"""
    pie = PieChart()
    pie.title = f"预算分配占比（{bo.total_a}万）"
    pie.height = 14
    pie.width = 18
    pa_data = Reference(ws, min_col=3, min_row=4, max_row=3 + len(bo.rows))
    pa_cats = Reference(ws, min_col=2, min_row=4, max_row=3 + len(bo.rows))
    pie.add_data(pa_data, titles_from_data=False)
    pie.set_categories(pa_cats)
    pie.dataLabels = DataLabelList()
    pie.dataLabels.showPercent = True
    pie.dataLabels.showCatName = True
    ws.add_chart(pie, f"A{total_row + 3}")


def _add_charts_sheet1(ws, bo, va, vb, total_row) -> None:
    """为预算总览对比表添加图表（柱状图 + 饼图）。"""
    # 柱状图 - 预算对比
    bar = BarChart()
    bar.type = "col"
    bar.style = 10
    bar.title = f"预算分配对比（{va}" + (f" vs {vb}" if vb else "）")
    bar.y_axis.title = "金额（万元）"
    bar.x_axis.title = "费用项目"
    bar.height = 12
    bar.width = 22

    cats = Reference(ws, min_col=2, min_row=4, max_row=3 + len(bo.rows))
    data_refs = [
        (3, "4A6CF7"),
    ]
    if vb:
        data_refs.append((5, "F59E0B"))

    for col_idx, color in data_refs:
        d = Reference(
            ws, min_col=col_idx, min_row=3,
            max_col=col_idx, max_row=3 + len(bo.rows),
        )
        bar.add_data(d, titles_from_data=True)
    bar.set_categories(cats)
    for i in range(len(data_refs)):
        bar.series[i].graphicalProperties.solidFill = data_refs[i][1]
    bar.gapWidth = 80
    ws.add_chart(bar, f"A{total_row + 3}")

    # 饼图 - A版本
    pie_a = PieChart()
    pie_a.title = f"{va}预算分配占比（{bo.total_a}万）"
    pie_a.height = 12
    pie_a.width = 16
    pa_data = Reference(ws, min_col=3, min_row=4, max_row=3 + len(bo.rows))
    pa_cats = Reference(ws, min_col=2, min_row=4, max_row=3 + len(bo.rows))
    pie_a.add_data(pa_data, titles_from_data=False)
    pie_a.set_categories(pa_cats)
    pie_a.dataLabels = DataLabelList()
    pie_a.dataLabels.showPercent = True
    pie_a.dataLabels.showCatName = True
    ws.add_chart(pie_a, f"A{total_row + 19}")

    # B饼图
    if vb:
        pie_b = PieChart()
        pie_b.title = f"{vb}预算分配占比（{bo.total_b}万）"
        pie_b.height = 12
        pie_b.width = 16
        pb_data = Reference(ws, min_col=5, min_row=4, max_row=3 + len(bo.rows))
        pb_cats = Reference(ws, min_col=2, min_row=4, max_row=3 + len(bo.rows))
        pie_b.add_data(pb_data, titles_from_data=False)
        pie_b.set_categories(pb_cats)
        pie_b.dataLabels = DataLabelList()
        pie_b.dataLabels.showPercent = True
        pie_b.dataLabels.showCatName = True
        ws.add_chart(pie_b, f"A{total_row + 35}")


# ── Sheet 2: 预算明细 ──────────────────────────────────────


def _build_sheet2(wb, data: XlsxData) -> None:
    bi = data.brand_info
    detail = data.budget_detail
    ws = wb.create_sheet("预算明细")

    ws.merge_cells("A1:F1")
    _apply_cell(
        ws["A1"], "预算明细 — 各项费用拆分",
        TITLE_FONT, TITLE_FILL, CENTER,
    )
    ws.row_dimensions[1].height = 36

    va, vb = bi.version_a_name, bi.version_b_name
    headers = ["费用大类", "子项明细", "用途说明", f"{va}(万元)"]
    if vb:
        headers += [f"{vb}(万元)"]
    headers.append("备注")

    for i, h in enumerate(headers, 1):
        _apply_cell(ws.cell(row=2, column=i), h, HEAD_FONT, HEAD_FILL)
    ws.row_dimensions[2].height = 28

    for idx, item in enumerate(detail.items, 3):
        row_vals = [item.category, item.sub, item.desc, item.amount_a]
        if vb:
            row_vals += [item.amount_b]
        row_vals.append(item.note)

        for c, val in enumerate(row_vals, 1):
            align = LEFT if c in (2, 3, 6 if vb else 5) else CENTER
            num_fmt = "0.0" if c in (4, 5 if vb else 4) else None
            _apply_cell(
                ws.cell(row=idx, column=c), val,
                CELL_FONT, align=align,
                num_fmt=num_fmt,
            )
        # 大类列着色
        ws.cell(row=idx, column=1).fill = _get_cat_fill(item.category)
        ws.cell(row=idx, column=1).font = SUBHEAD_FONT
        ws.row_dimensions[idx].height = 24

    # 合计行
    trow = 3 + len(detail.items)
    sum_a = sum(i.amount_a for i in detail.items)
    sum_b = sum((i.amount_b or 0) for i in detail.items)
    total_vals = ["合计", "", "", sum_a]
    if vb:
        total_vals += [sum_b]
    total_vals.append("")
    for c, val in enumerate(total_vals, 1):
        _apply_cell(
            ws.cell(row=trow, column=c), val,
            TOTAL_FONT, TOTAL_FILL, num_fmt="0.0" if c in (4, 5) else None,
        )
    ws.row_dimensions[trow].height = 30

    col_widths = [26, 24, 40, 14]
    if vb:
        col_widths += [14]
    col_widths.append(22)
    _set_widths(ws, col_widths)

    # 分类汇总
    _add_category_summary(ws, detail, va, vb, trow)


def _add_category_summary(ws, detail, va, vb, trow) -> None:
    """添加费用大类汇总数据 + 柱状图。"""
    cat_start = trow + 3
    headers = ["费用大类", f"{va}(万)"]
    if vb:
        headers.append(f"{vb}(万)")
    for i, h in enumerate(headers, 1):
        _apply_cell(ws.cell(row=cat_start, column=i), h, HEAD_FONT, HEAD_FILL)

    cat_sum: dict[str, list[float]] = {}
    for item in detail.items:
        if item.category not in cat_sum:
            cat_sum[item.category] = [0.0, 0.0]
        cat_sum[item.category][0] += item.amount_a
        cat_sum[item.category][1] += item.amount_b or 0

    for i, (cat, (v1, v2)) in enumerate(cat_sum.items()):
        r = cat_start + 1 + i
        _apply_cell(ws.cell(row=r, column=1), cat, CELL_FONT, align=LEFT)
        _apply_cell(ws.cell(row=r, column=2), v1, CELL_FONT, num_fmt="0.0")
        if vb:
            _apply_cell(ws.cell(row=r, column=3), v2, CELL_FONT, num_fmt="0.0")

    cat_count = len(cat_sum)
    if cat_count < 2:
        return

    bar_cat = BarChart()
    bar_cat.type = "col"
    bar_cat.style = 10
    bar_cat.title = "各费用大类对比"
    bar_cat.y_axis.title = "金额（万元）"
    bar_cat.height = 12
    bar_cat.width = 22

    c_cats = Reference(ws, min_col=1, min_row=cat_start + 1, max_row=cat_start + cat_count)
    d1 = Reference(ws, min_col=2, min_row=cat_start, max_col=2, max_row=cat_start + cat_count)
    bar_cat.add_data(d1, titles_from_data=True)
    colors = ["4A6CF7"]
    if vb:
        d2 = Reference(ws, min_col=3, min_row=cat_start, max_col=3, max_row=cat_start + cat_count)
        bar_cat.add_data(d2, titles_from_data=True)
        colors.append("F59E0B")
    bar_cat.set_categories(c_cats)
    for i in range(len(colors)):
        bar_cat.series[i].graphicalProperties.solidFill = colors[i]
    bar_cat.gapWidth = 80
    ws.add_chart(bar_cat, f"E{cat_start}")


# ── Sheet 3: 活动流程时间线 ────────────────────────────────


def _build_sheet3(wb, data: XlsxData) -> None:
    bi = data.brand_info
    timeline = data.timeline
    ws = wb.create_sheet("活动流程时间线")

    ws.merge_cells("A1:G1")
    _apply_cell(
        ws["A1"],
        f"活动流程 — {bi.duration_months}个月执行时间线",
        TITLE_FONT, TITLE_FILL, CENTER,
    )
    ws.row_dimensions[1].height = 36

    headers = ["阶段", "时间", "阶段目标", "核心任务", "具体动作", "负责方", "阶段产出"]
    for i, h in enumerate(headers, 1):
        _apply_cell(ws.cell(row=2, column=i), h, HEAD_FONT, HEAD_FILL)
    ws.row_dimensions[2].height = 28

    for idx, item in enumerate(timeline.items, 3):
        phase_fill = PHASE_FILLS.get(item.phase, GREY_FILL)
        row_vals = [
            item.phase, item.week, item.goal, item.task,
            item.action, item.owner, item.deliverable,
        ]
        for c, val in enumerate(row_vals, 1):
            align = LEFT if c >= 4 else CENTER
            font = SUBHEAD_FONT if c == 1 else CELL_FONT
            cell_fill = phase_fill if c == 1 else None
            _apply_cell(
                ws.cell(row=idx, column=c), val,
                font, cell_fill, align,
            )
        ws.row_dimensions[idx].height = 42

    _set_widths(ws, [10, 10, 18, 22, 48, 18, 20])

    # 甘特图
    _add_gantt_chart(ws, timeline)


def _add_gantt_chart(ws, timeline) -> None:
    """添加甘特图阶段汇总 + 条形图。"""
    gantt_start = 3 + len(timeline.items) + 3
    _apply_cell(
        ws.cell(row=gantt_start, column=1), "阶段", HEAD_FONT, HEAD_FILL,
    )
    _apply_cell(
        ws.cell(row=gantt_start, column=2), "起始周", HEAD_FONT, HEAD_FILL,
    )
    _apply_cell(
        ws.cell(row=gantt_start, column=3), "持续周数", HEAD_FONT, HEAD_FILL,
    )

    for i, item in enumerate(timeline.gantt):
        r = gantt_start + 1 + i
        _apply_cell(ws.cell(row=r, column=1), item.phase, CELL_FONT)
        _apply_cell(ws.cell(row=r, column=2), item.start_week, CELL_FONT)
        _apply_cell(ws.cell(row=r, column=3), item.duration_weeks, CELL_FONT)

    if not timeline.gantt:
        return

    gantt_bar = BarChart()
    gantt_bar.type = "bar"
    gantt_bar.style = 12
    gantt_bar.title = "活动流程甘特图"
    gantt_bar.x_axis.title = "周次"
    gantt_bar.height = 8
    gantt_bar.width = 20
    g_cats = Reference(
        ws, min_col=1, min_row=gantt_start + 1,
        max_row=gantt_start + len(timeline.gantt),
    )
    g_data = Reference(
        ws, min_col=3, min_row=gantt_start,
        max_col=3, max_row=gantt_start + len(timeline.gantt),
    )
    gantt_bar.add_data(g_data, titles_from_data=True)
    gantt_bar.set_categories(g_cats)
    gantt_bar.series[0].graphicalProperties.solidFill = "4A6CF7"
    gantt_bar.gapWidth = 50
    ws.add_chart(gantt_bar, f"E{gantt_start}")


# ── Sheet 4: 预期效果KPI ───────────────────────────────────


def _build_sheet4(wb, data: XlsxData) -> None:
    bi = data.brand_info
    kpi_data = data.kpi
    ws = wb.create_sheet("预期效果KPI")

    ws.merge_cells("A1:F1")
    _apply_cell(
        ws["A1"], "预期效果与 KPI — 四维指标体系",
        TITLE_FONT, TITLE_FILL, CENTER,
    )
    ws.row_dimensions[1].height = 36

    va, vb = bi.version_a_name, bi.version_b_name
    headers = ["指标维度", "指标名称", "指标定义", f"{va}目标值"]
    if vb:
        headers += [f"{vb}目标值"]
    headers.append("达成路径")

    for i, h in enumerate(headers, 1):
        _apply_cell(ws.cell(row=2, column=i), h, HEAD_FONT, HEAD_FILL)
    ws.row_dimensions[2].height = 28

    col_count = len(headers)
    for idx, kpi_row in enumerate(kpi_data.kpis, 3):
        dim_fill = DIM_FILLS.get(kpi_row.dim, GREY_FILL)
        row_vals = [
            kpi_row.dim, kpi_row.metric, kpi_row.definition,
            kpi_row.target_a,
        ]
        if vb:
            row_vals += [kpi_row.target_b]
        row_vals.append(kpi_row.path)

        for c, val in enumerate(row_vals, 1):
            align = LEFT if c in (2, 3, col_count) else CENTER
            font = SUBHEAD_FONT if c == 1 else CELL_FONT
            cell_fill = dim_fill if c == 1 else None
            _apply_cell(
                ws.cell(row=idx, column=c), val,
                font, cell_fill, align,
            )
        ws.row_dimensions[idx].height = 30

    # 辅助数值列
    num_col_start = col_count + 2
    _apply_cell(
        ws.cell(row=2, column=num_col_start), f"{va}数值",
        HEAD_FONT, HEAD_FILL,
    )
    ws.column_dimensions[get_column_letter(num_col_start)].width = 14
    if vb:
        _apply_cell(
            ws.cell(row=2, column=num_col_start + 1), f"{vb}数值",
            HEAD_FONT, HEAD_FILL,
        )
        ws.column_dimensions[get_column_letter(num_col_start + 1)].width = 14

    for i, kpi_row in enumerate(kpi_data.kpis):
        r = 3 + i
        _apply_cell(
            ws.cell(row=r, column=num_col_start), kpi_row.target_a_num,
            CELL_FONT, num_fmt="0.0",
        )
        if vb and kpi_row.target_b_num is not None:
            _apply_cell(
                ws.cell(row=r, column=num_col_start + 1),
                kpi_row.target_b_num,
                CELL_FONT, num_fmt="0.0",
            )

    # ROI 分析
    roi_start = 3 + len(kpi_data.kpis) + 1
    ws.merge_cells(f"A{roi_start}:{get_column_letter(col_count)}{roi_start}")
    _apply_cell(
        ws.cell(row=roi_start, column=1), "ROI 投入产出分析",
        SUBHEAD_FONT, SUBHEAD_FILL, CENTER,
    )
    ws.row_dimensions[roi_start].height = 28

    for i, roi_row in enumerate(kpi_data.roi):
        r = roi_start + 1 + i
        row_vals = [
            roi_row.type, roi_row.metric, roi_row.definition,
            roi_row.value_a,
        ]
        if vb:
            row_vals += [roi_row.value_b]
        row_vals.append(roi_row.note)
        for c, val in enumerate(row_vals, 1):
            align = LEFT if c in (2, 3, col_count) else CENTER
            font = SUBHEAD_FONT if c == 1 else CELL_FONT
            _apply_cell(
                ws.cell(row=r, column=c), val,
                font, TOTAL_FILL if c == 1 else None, align,
            )
        ws.row_dimensions[r].height = 28

    _set_widths(ws, [12, 22, 36, 18, 18, 30])

    # 图表
    _add_kpi_charts(ws, kpi_data, va, vb, roi_start, col_count, num_col_start)


def _add_kpi_charts(ws, kpi_data, va, vb, roi_start, col_count, num_col_start) -> None:
    """添加 KPI 对比柱状图、ROI 折线图、四维雷达图。"""
    kpi_count = min(8, len(kpi_data.kpis))
    chart_start = roi_start + 8

    # KPI 对比柱状图
    if kpi_count > 0:
        bar_kpi = BarChart()
        bar_kpi.type = "bar"
        bar_kpi.style = 11
        bar_kpi.title = f"核心 KPI 对比"
        bar_kpi.x_axis.title = "数值"
        bar_kpi.y_axis.title = "指标"
        bar_kpi.height = 14
        bar_kpi.width = 24
        k_cats = Reference(ws, min_col=2, min_row=3, max_row=2 + kpi_count)
        k_d1 = Reference(
            ws, min_col=num_col_start, min_row=2,
            max_col=num_col_start, max_row=2 + kpi_count,
        )
        bar_kpi.add_data(k_d1, titles_from_data=True)
        colors = ["4A6CF7"]
        if vb:
            k_d2 = Reference(
                ws, min_col=num_col_start + 1, min_row=2,
                max_col=num_col_start + 1, max_row=2 + kpi_count,
            )
            bar_kpi.add_data(k_d2, titles_from_data=True)
            colors.append("10B981")
        bar_kpi.set_categories(k_cats)
        for i in range(len(colors)):
            bar_kpi.series[i].graphicalProperties.solidFill = colors[i]
        bar_kpi.gapWidth = 80
        ws.add_chart(bar_kpi, f"A{chart_start}")

    # ROI 折线图
    if kpi_data.roi_chart:
        roi_chart_start = chart_start + 25
        rc_headers = ["项目", va]
        if vb:
            rc_headers.append(vb)
        for i, h in enumerate(rc_headers, 1):
            _apply_cell(
                ws.cell(row=roi_chart_start, column=i), h,
                HEAD_FONT, HEAD_FILL,
            )
        for i, pt in enumerate(kpi_data.roi_chart):
            r = roi_chart_start + 1 + i
            _apply_cell(
                ws.cell(row=r, column=1), pt.name, CELL_FONT, align=LEFT,
            )
            _apply_cell(
                ws.cell(row=r, column=2), pt.value_a, CELL_FONT, num_fmt="0",
            )
            if vb and pt.value_b is not None:
                _apply_cell(
                    ws.cell(row=r, column=3), pt.value_b, CELL_FONT,
                    num_fmt="0",
                )

        line_roi = LineChart()
        line_roi.title = "投入产出对比（预算 vs 预估营收）"
        line_roi.style = 12
        line_roi.y_axis.title = "金额（万元）"
        line_roi.height = 10
        line_roi.width = 18
        r_cats = Reference(
            ws, min_col=1, min_row=roi_chart_start + 1,
            max_row=roi_chart_start + len(kpi_data.roi_chart),
        )
        r_d1 = Reference(
            ws, min_col=2, min_row=roi_chart_start,
            max_col=2, max_row=roi_chart_start + len(kpi_data.roi_chart),
        )
        line_roi.add_data(r_d1, titles_from_data=True)
        line_roi.series[0].graphicalProperties.line.solidFill = "4A6CF7"
        line_roi.series[0].graphicalProperties.line.width = 30000
        if vb:
            r_d2 = Reference(
                ws, min_col=3, min_row=roi_chart_start,
                max_col=3, max_row=roi_chart_start + len(kpi_data.roi_chart),
            )
            line_roi.add_data(r_d2, titles_from_data=True)
            line_roi.series[1].graphicalProperties.line.solidFill = "F59E0B"
            line_roi.series[1].graphicalProperties.line.width = 30000
        ws.add_chart(line_roi, f"E{chart_start}")

    # 雷达图
    if kpi_data.radar_chart:
        radar_start = (roi_chart_start if kpi_data.roi_chart else chart_start + 25) + len(kpi_data.roi_chart) + 3
        _apply_cell(
            ws.cell(row=radar_start, column=1), "维度",
            HEAD_FONT, HEAD_FILL,
        )
        _apply_cell(
            ws.cell(row=radar_start, column=2), f"{va}得分",
            HEAD_FONT, HEAD_FILL,
        )
        if vb:
            _apply_cell(
                ws.cell(row=radar_start, column=3), f"{vb}得分",
                HEAD_FONT, HEAD_FILL,
            )
        for i, sc in enumerate(kpi_data.radar_chart):
            r = radar_start + 1 + i
            _apply_cell(
                ws.cell(row=r, column=1), sc.dim, CELL_FONT, align=LEFT,
            )
            _apply_cell(
                ws.cell(row=r, column=2), sc.score_a, CELL_FONT,
            )
            if vb and sc.score_b is not None:
                _apply_cell(
                    ws.cell(row=r, column=3), sc.score_b, CELL_FONT,
                )

        radar = RadarChart()
        radar.type = "filled"
        radar.style = 26
        radar.title = "四维营销效果雷达图"
        radar.height = 12
        radar.width = 16
        rd_cats = Reference(
            ws, min_col=1, min_row=radar_start + 1,
            max_row=radar_start + len(kpi_data.radar_chart),
        )
        rd_d1 = Reference(
            ws, min_col=2, min_row=radar_start,
            max_col=2, max_row=radar_start + len(kpi_data.radar_chart),
        )
        radar.add_data(rd_d1, titles_from_data=True)
        radar.set_categories(rd_cats)
        if vb:
            rd_d2 = Reference(
                ws, min_col=3, min_row=radar_start,
                max_col=3, max_row=radar_start + len(kpi_data.radar_chart),
            )
            radar.add_data(rd_d2, titles_from_data=True)
        ws.add_chart(radar, f"E{radar_start + 5}")


# ── Tool ──


@tool(args_schema=GeneratePlanXlsxInput)
async def generate_plan_xlsx(data: str, brand_name: str) -> str:
    """根据营销方案结构化数据生成 XLSX 预算流程回报分析表格。

    数据由 LLM with_structured_output 提取自方案流水线上游的各节点输出，
    返回生成文件的绝对路径。
    """
    try:
        parsed = json.loads(data) if isinstance(data, str) else data
        xlsx_data = XlsxData.model_validate(parsed)
    except Exception as exc:
        logger.exception("generate_plan_xlsx: Invalid XlsxData")
        return f"数据解析失败：{exc}"

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    today = date.today().strftime("%Y%m%d")
    filename = f"{xlsx_data.brand_info.brand_name}_{today}_预算流程回报分析.xlsx"
    filepath = OUTPUT_DIR / filename

    # ponytail: 如果明细汇总远超总预算，很可能 LLM 输出了重复的汇总行，需提示词修正
    detail_sum = sum(i.amount_a for i in xlsx_data.budget_detail.items)
    budget = xlsx_data.brand_info.version_a_budget
    if detail_sum > budget * 1.3:
        logger.warning(
            "budget detail sum (%.1f) >> total budget (%.1f, ratio=%.2f), "
            "likely LLM included summary rows. Check prompt xlsx_generation.md.j2",
            detail_sum, budget, detail_sum / budget,
        )

    try:
        wb = openpyxl.Workbook()
        _build_sheet1(wb, xlsx_data)
        _build_sheet2(wb, xlsx_data)
        _build_sheet3(wb, xlsx_data)
        _build_sheet4(wb, xlsx_data)

        wb.save(str(filepath))
        abs_path = str(filepath.resolve())
        logger.info("generate_plan_xlsx: saved to %s", abs_path)
        return abs_path
    except Exception as exc:
        logger.exception("generate_plan_xlsx: save failed")
        return f"XLSX 生成失败：{exc}"
